from datetime import datetime, timedelta
import io

import openpyxl
from bson import ObjectId


def _create_survey(client, headers, settings=None, deadline=None):
    payload = {
        'title': '第一阶段验收问卷',
        'description': '自动化验收',
        'settings': settings if settings is not None else {},
        'deadline': deadline,
    }
    res = client.post('/api/surveys', json=payload, headers=headers)
    assert res.status_code == 201, res.get_json()
    return res.get_json()['data']


def _add_question(client, headers, survey_id, title, qtype='text', order=0, required=True, **extra):
    payload = {'title': title, 'type': qtype, 'order': order, 'required': required}
    payload.update(extra)
    res = client.post(f'/api/surveys/{survey_id}/questions', json=payload, headers=headers)
    assert res.status_code == 201, res.get_json()
    return res.get_json()['data']['id']


def _start_anonymous(client, survey_id):
    return client.post(f'/api/fill/surveys/{survey_id}/start', json={})


def _submit(client, survey_id, answers, response_id, headers=None):
    return client.post(f'/api/fill/surveys/{survey_id}/submit', json={
        'answers': answers,
        'response_id': response_id,
    }, headers=headers)


def test_anonymous_survey_blocks_anonymous_start_and_submit(client, auth_headers, db):
    survey = _create_survey(client, auth_headers, {'allow_anonymous': False})
    qid = _add_question(client, auth_headers, survey['id'], '必填题')
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)

    assert _start_anonymous(client, survey['id']).status_code == 400
    res = client.post(f"/api/fill/surveys/{survey['id']}/start", json={}, headers=auth_headers)
    assert res.status_code == 201
    rid = res.get_json()['data']['response_id']
    submit_res = _submit(client, survey['id'], {qid: 'y'}, rid, headers=auth_headers)
    assert submit_res.status_code == 200, submit_res.get_json()

    # 已开始的匿名回答，在设置改为“仅登录”后也不能提交
    survey2 = _create_survey(client, auth_headers, {'allow_anonymous': True})
    _add_question(client, auth_headers, survey2['id'], '必填题')
    client.post(f"/api/surveys/{survey2['id']}/publish", headers=auth_headers)
    rid2 = _start_anonymous(client, survey2['id']).get_json()['data']['response_id']
    db.surveys.update_one({'_id': ObjectId(survey2['id'])}, {'$set': {'settings.allow_anonymous': False}})
    assert _submit(client, survey2['id'], {}, rid2).status_code == 400


def test_deadline_blocks_start_submit_and_link(client, auth_headers, db):
    future = (datetime.utcnow() + timedelta(days=1)).isoformat() + 'Z'
    survey = _create_survey(client, auth_headers, {'allow_anonymous': True}, deadline=future)
    _add_question(client, auth_headers, survey['id'], '必填题')
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)

    link_res = client.get(f"/api/surveys/link/{survey['link_code']}")
    assert link_res.status_code == 200
    assert link_res.get_json()['data']['title'] == survey['title']

    rid = _start_anonymous(client, survey['id']).get_json()['data']['response_id']
    past = datetime.utcnow() - timedelta(minutes=1)
    db.surveys.update_one({'_id': ObjectId(survey['id'])}, {'$set': {'deadline': past}})

    assert _start_anonymous(client, survey['id']).status_code == 400
    assert _submit(client, survey['id'], {}, rid).status_code == 400
    assert client.get(f"/api/surveys/link/{survey['link_code']}").status_code == 404



def test_closed_survey_is_locked(client, auth_headers):
    survey = _create_survey(client, auth_headers)
    qid = _add_question(client, auth_headers, survey['id'], '题目')
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)
    assert client.post(f"/api/surveys/{survey['id']}/close", headers=auth_headers).status_code == 200

    assert client.put(f"/api/surveys/{survey['id']}", json={'title': '改标题'}, headers=auth_headers).status_code == 400
    assert client.post(f"/api/surveys/{survey['id']}/questions", json={
        'title': '新题', 'type': 'text', 'order': 1
    }, headers=auth_headers).status_code == 400
    assert client.put(f'/api/surveys/questions/{qid}', json={'title': '改题目'},
                      headers=auth_headers).status_code == 400
    assert client.delete(f'/api/surveys/questions/{qid}', headers=auth_headers).status_code == 400


def test_stats_counts_only_submitted_responses(client, auth_headers):
    survey = _create_survey(client, auth_headers, {'allow_anonymous': True})
    qid = _add_question(client, auth_headers, survey['id'], '你的建议')
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)

    rid1 = _start_anonymous(client, survey['id']).get_json()['data']['response_id']
    rid2 = _start_anonymous(client, survey['id']).get_json()['data']['response_id']
    assert rid1 != rid2
    assert _submit(client, survey['id'], {qid: '很好'}, rid1).status_code == 200

    res = client.get(f"/api/stats/surveys/{survey['id']}", headers=auth_headers)
    assert res.status_code == 200
    data = res.get_json()['data']
    assert data['total_responses'] == 1
    assert data['completion_rate'] == 0.5
    assert data['question_stats'][0]['title'] == '你的建议'
    assert data['question_stats'][0]['total_answers'] == 1


def test_hide_jump_skips_required_question_and_discards_stale_answer(client, auth_headers, db):
    survey = _create_survey(client, auth_headers, {'allow_anonymous': True})
    q0 = _add_question(client, auth_headers, survey['id'], '是否同意', 'single_choice', 0,
                       options=['同意', '不同意'])
    q1 = _add_question(client, auth_headers, survey['id'], '补充说明', 'text', 1)
    q2 = _add_question(client, auth_headers, survey['id'], '联系方式', 'text', 2)
    jump = {
        'source_question_id': q0,
        'condition': 'equals',
        'value': '同意',
        'target_question_id': q1,
        'action': 'hide',
    }
    assert client.put(f'/api/surveys/questions/{q1}', json={'jump_logic': jump},
                      headers=auth_headers).status_code == 200
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)

    rid = _start_anonymous(client, survey['id']).get_json()['data']['response_id']
    res = _submit(client, survey['id'], {q0: '同意', q1: '不应保存', q2: '123'}, rid)
    assert res.status_code == 200, res.get_json()

    stored = db.responses.find_one({'_id': ObjectId(rid)})
    assert q0 in stored['answers']
    assert q1 not in stored['answers']
    assert q2 in stored['answers']


def test_skip_to_jump_skips_intermediate_required_question(client, auth_headers):
    survey = _create_survey(client, auth_headers, {'allow_anonymous': True})
    q0 = _add_question(client, auth_headers, survey['id'], '选择路径', 'single_choice', 0,
                       options=['直达', '普通'])
    q2 = _add_question(client, auth_headers, survey['id'], '目标题', 'text', 2)
    _add_question(client, auth_headers, survey['id'], '中间题', 'text', 1)
    jump = {
        'source_question_id': q0,
        'condition': 'equals',
        'value': '直达',
        'target_question_id': q2,
        'action': 'skip_to',
    }
    assert client.put(f'/api/surveys/questions/{q0}', json={'jump_logic': jump},
                      headers=auth_headers).status_code == 200
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)

    rid = _start_anonymous(client, survey['id']).get_json()['data']['response_id']
    res = _submit(client, survey['id'], {q0: '直达', q2: '已填写'}, rid)
    assert res.status_code == 200, res.get_json()


def test_required_multiple_choice_rejects_empty_list_and_duplicate_submit(client, auth_headers):
    survey = _create_survey(client, auth_headers, {'allow_anonymous': True})
    qid = _add_question(client, auth_headers, survey['id'], '多选必填', 'multiple_choice', 0,
                        options=['A', 'B'])
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)
    rid = _start_anonymous(client, survey['id']).get_json()['data']['response_id']

    empty_res = _submit(client, survey['id'], {qid: []}, rid)
    assert empty_res.status_code == 422
    assert qid in empty_res.get_json()['data']

    ok_res = _submit(client, survey['id'], {qid: ['A']}, rid)
    assert ok_res.status_code == 200, ok_res.get_json()
    assert _submit(client, survey['id'], {qid: ['A']}, rid).status_code == 400


def test_collect_email_is_validated_and_exported(client, auth_headers):
    survey = _create_survey(client, auth_headers, {
        'allow_anonymous': True,
        'collect_email': True,
    })
    qid = _add_question(client, auth_headers, survey['id'], '邮箱收集题')
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)
    rid = _start_anonymous(client, survey['id']).get_json()['data']['response_id']

    invalid = client.post(f"/api/fill/surveys/{survey['id']}/submit", json={
        'answers': {qid: '答案'},
        'response_id': rid,
        'email': 'not-an-email',
    })
    assert invalid.status_code == 422
    assert 'email' in invalid.get_json()['data']

    valid = client.post(f"/api/fill/surveys/{survey['id']}/submit", json={
        'answers': {qid: '答案'},
        'response_id': rid,
        'email': 'fill@example.com',
    })
    assert valid.status_code == 200, valid.get_json()

    export_res = client.get(f"/api/stats/surveys/{survey['id']}/export?format=json",
                            headers=auth_headers)
    assert export_res.status_code == 200
    assert export_res.get_json()['data'][0]['email'] == 'fill@example.com'

    csv_res = client.get(f"/api/stats/surveys/{survey['id']}/export?format=csv",
                         headers=auth_headers)
    assert csv_res.status_code == 200
    assert csv_res.data.decode('utf-8-sig').splitlines()[0].startswith(
        'response_id,user_id,email,submitted_at')
    assert 'fill@example.com' in csv_res.data.decode('utf-8-sig')

    xlsx_res = client.get(f"/api/stats/surveys/{survey['id']}/export?format=xlsx",
                          headers=auth_headers)
    assert xlsx_res.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_res.data))
    sheet = workbook.active
    assert sheet.cell(row=1, column=3).value == 'email'
    assert sheet.cell(row=2, column=3).value == 'fill@example.com'


def test_response_progress_and_draft_are_owner_only(client, auth_headers):
    survey = _create_survey(client, auth_headers, {'allow_anonymous': False})
    _add_question(client, auth_headers, survey['id'], '题目')
    client.post(f"/api/surveys/{survey['id']}/publish", headers=auth_headers)
    rid = client.post(f"/api/fill/surveys/{survey['id']}/start", json={},
                      headers=auth_headers).get_json()['data']['response_id']

    client.post('/api/auth/register', json={
        'username': 'other_user',
        'email': 'other@example.com',
        'password': 'OtherPass123',
    })
    login = client.post('/api/auth/login', json={
        'username': 'other_user',
        'password': 'OtherPass123',
    })
    other_headers = {'Authorization': f"Bearer {login.get_json()['data']['access_token']}"}

    assert client.get(f'/api/fill/responses/{rid}', headers=other_headers).status_code == 403
    assert client.put(f'/api/fill/responses/{rid}/draft', json={'answers': {}},
                      headers=other_headers).status_code == 400
    assert client.get(f'/api/fill/responses/{rid}', headers=auth_headers).status_code == 200

