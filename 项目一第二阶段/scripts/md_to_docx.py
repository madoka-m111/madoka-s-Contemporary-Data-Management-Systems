#!/usr/bin/env python3
"""将 docs/项目完成报告.md 转换为排版后的 Word 文档。

用法：
    python scripts/md_to_docx.py
    python scripts/md_to_docx.py docs/项目完成报告.md docs/项目完成报告.docx

依赖（仅文档转换时需要）：
    python -m pip install python-docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BODY_FONT = "宋体"
HEAD_FONT = "黑体"
CODE_FONT = "Consolas"
TEXT_COLOR = RGBColor(0x1F, 0x29, 0x37)
MUTED_COLOR = RGBColor(0x59, 0x59, 0x59)
HEADER_COLOR = RGBColor(0x1F, 0x38, 0x64)
CODE_COLOR = RGBColor(0x1F, 0x4E, 0x79)

CONTENT_WIDTH_CM = 15.0


# ---------------------------------------------------------------------------
# 基础 XML / 字体辅助函数
# ---------------------------------------------------------------------------

def _set_rfonts(rpr, ascii_font: str, east_asia: str) -> None:
    """显式设置字体，并移除会覆盖显式字体的主题字体属性。"""
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), ascii_font)


def set_style_font(
    style,
    ascii_font: str = BODY_FONT,
    east_asia: str = BODY_FONT,
    size: float = 10.5,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    style.font.name = ascii_font
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    _set_rfonts(style.element.get_or_add_rPr(), ascii_font, east_asia)


def set_run_font(
    run,
    ascii_font: str = BODY_FONT,
    east_asia: str = BODY_FONT,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    _set_rfonts(rpr, ascii_font, east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_run(run, fill: str) -> None:
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rpr.append(shd)


def shade_paragraph(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_paragraph_border(
    paragraph,
    edge: str = "bottom",
    color: str = "D9D9D9",
    size: int = 6,
    space: int = 2,
) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.first_child_found_in("w:pBdr")
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    pbdr.append(element)


def add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    """插入可在 Word 中自动更新的域（PAGE / NUMPAGES）。"""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f" {instruction} ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "595959")
    rpr.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    rpr.append(size)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def set_update_fields(doc) -> None:
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def set_cell_margins(cell, top: int = 80, bottom: int = 80,
                     left: int = 110, right: int = 110) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for tag, value in (("top", top), ("bottom", bottom),
                       ("start", left), ("end", right)):
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        tcmar.append(element)


def set_table_borders(table, color: str = "BFBFBF", size: int = 4) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tblpr.append(borders)


def set_table_width(table) -> None:
    tblpr = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    tblpr.append(width)


def repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


# ---------------------------------------------------------------------------
# 文档样式与页面设置
# ---------------------------------------------------------------------------

def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_style_font(normal, size=10.5, color=TEXT_COLOR)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(21)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    heading_specs = {
        "Heading 1": (16, 18, 8),
        "Heading 2": (14, 12, 6),
        "Heading 3": (12, 8, 4),
    }
    for name, (size, before, after) in heading_specs.items():
        style = doc.styles[name]
        set_style_font(style, ascii_font=HEAD_FONT, east_asia=HEAD_FONT,
                       size=size, bold=True, color=TEXT_COLOR)
        hpf = style.paragraph_format
        hpf.first_line_indent = Pt(0)
        hpf.space_before = Pt(before)
        hpf.space_after = Pt(after)
        hpf.line_spacing = 1.2
        hpf.keep_with_next = True
        hpf.keep_together = True
        hpf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def add_style(name: str):
        try:
            return doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = doc.styles["Normal"]
            style.next_paragraph_style = doc.styles["Normal"]
            return style

    title = add_style("ReportTitle")
    set_style_font(title, ascii_font=HEAD_FONT, east_asia=HEAD_FONT,
                   size=24, bold=True, color=TEXT_COLOR)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_before = Pt(96)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.line_spacing = 1.2

    subtitle = add_style("ReportSubtitle")
    set_style_font(subtitle, ascii_font=HEAD_FONT, east_asia=HEAD_FONT,
                   size=16, bold=True, color=TEXT_COLOR)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(48)
    subtitle.paragraph_format.line_spacing = 1.2

    cover_name = add_style("CoverName")
    set_style_font(cover_name, size=12, color=TEXT_COLOR)
    cover_name.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_name.paragraph_format.first_line_indent = Pt(0)
    cover_name.paragraph_format.space_before = Pt(36)
    cover_name.paragraph_format.space_after = Pt(0)

    code = add_style("CodeBlock")
    set_style_font(code, ascii_font=CODE_FONT, east_asia=BODY_FONT,
                   size=8.5, color=TEXT_COLOR)
    code.paragraph_format.first_line_indent = Pt(0)
    code.paragraph_format.left_indent = Cm(0.35)
    code.paragraph_format.right_indent = Cm(0.2)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.keep_together = True

    table_text = add_style("TableText")
    set_style_font(table_text, size=9, color=TEXT_COLOR)
    table_text.paragraph_format.first_line_indent = Pt(0)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.15
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    list_text = add_style("ListText")
    set_style_font(list_text, size=10.5, color=TEXT_COLOR)
    list_text.paragraph_format.first_line_indent = Pt(0)
    list_text.paragraph_format.space_before = Pt(0)
    list_text.paragraph_format.space_after = Pt(2)
    list_text.paragraph_format.line_spacing = 1.4
    list_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("基于 MongoDB 的在线问卷系统")
    set_run_font(hr, size=8.5, color=MUTED_COLOR)
    add_paragraph_border(hp, "bottom", "D9D9D9", 4, 2)

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("第 ")
    set_run_font(fr, size=8.5, color=MUTED_COLOR)
    add_field(fp, "PAGE", "1")
    fr = fp.add_run(" 页 / 共 ")
    set_run_font(fr, size=8.5, color=MUTED_COLOR)
    add_field(fp, "NUMPAGES", "1")
    fr = fp.add_run(" 页")
    set_run_font(fr, size=8.5, color=MUTED_COLOR)

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    first_footer.paragraphs[0].text = ""

    set_update_fields(doc)


# ---------------------------------------------------------------------------
# 行内格式与正文块
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`\n]+`|\*[^*\n]+\*)")


def add_text_run(
    paragraph,
    text: str,
    ascii_font: str = BODY_FONT,
    east_asia: str = BODY_FONT,
    size: float = 10.5,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = TEXT_COLOR,
):
    run = paragraph.add_run(text)
    set_run_font(run, ascii_font=ascii_font, east_asia=east_asia,
                 size=size, bold=bold, italic=italic, color=color)
    if text[:1].isspace() or text[-1:].isspace():
        text_element = run._r.find(qn("w:t"))
        if text_element is not None:
            text_element.set(qn("xml:space"), "preserve")
    return run


def add_inline_text(
    paragraph,
    text: str,
    size: float = 10.5,
    color: RGBColor | None = TEXT_COLOR,
    bold: bool = False,
    italic: bool = False,
) -> None:
    """解析 **加粗**、`行内代码` 和 *斜体*。"""
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            add_text_run(paragraph, text[position:match.start()], size=size,
                         color=color, bold=bold, italic=italic)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            add_inline_text(paragraph, token[2:-2], size=size,
                            color=color, bold=True, italic=italic)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, ascii_font=CODE_FONT, east_asia=BODY_FONT,
                         size=max(8.0, size - 1.0), bold=bold,
                         italic=italic, color=CODE_COLOR)
            shade_run(run, "F2F2F2")
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            add_inline_text(paragraph, token[1:-1], size=size,
                            color=color, bold=bold, italic=True)
        else:
            add_text_run(paragraph, token, size=size,
                         color=color, bold=bold, italic=italic)
        position = match.end()
    if position < len(text):
        add_text_run(paragraph, text[position:], size=size,
                     color=color, bold=bold, italic=italic)


def add_body_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="Normal")
    add_inline_text(paragraph, text)
    return paragraph


def add_cover_name(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="CoverName")
    add_inline_text(paragraph, text, size=12)
    return paragraph


def add_heading_block(doc: Document, level: int, text: str):
    if level == 1:
        style_name, size = "ReportTitle", 24
    elif level == 2 and text.startswith("基于 MongoDB"):
        style_name, size = "ReportSubtitle", 16
    elif level == 2:
        style_name, size = "Heading 1", 16
    elif level == 3:
        style_name, size = "Heading 2", 14
    else:
        style_name, size = "Heading 3", 12

    paragraph = doc.add_paragraph(style=style_name)
    if re.match(r"^(摘要|\d+\.\s|附录)", text):
        paragraph.paragraph_format.page_break_before = True
    add_inline_text(paragraph, text, size=size, bold=True)
    return paragraph


def add_list_item(doc: Document, level: int, ordered: bool,
                  marker: str, text: str):
    paragraph = doc.add_paragraph(style="ListText")
    indent_cm = 0.74 * (level + 1)
    paragraph.paragraph_format.left_indent = Cm(indent_cm)
    paragraph.paragraph_format.first_line_indent = Cm(-0.74)
    paragraph.paragraph_format.space_after = Pt(1 if level else 2)

    if ordered:
        prefix = f"{marker}. "
    else:
        prefix = ("•", "◦", "▪")[min(level, 2)] + " "
    add_text_run(paragraph, prefix, size=10.5)
    add_inline_text(paragraph, text)
    return paragraph


def add_code_block(doc: Document, lines: list[str]):
    paragraph = doc.add_paragraph(style="CodeBlock")
    shade_paragraph(paragraph, "F5F5F5")
    for edge in ("top", "left", "bottom", "right"):
        add_paragraph_border(paragraph, edge, "D9D9D9", 4, 4)

    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, ascii_font=CODE_FONT, east_asia=BODY_FONT,
                     size=8.5, color=TEXT_COLOR)
        text_element = run._r.find(qn("w:t"))
        if text_element is not None:
            text_element.set(qn("xml:space"), "preserve")
    paragraph.paragraph_format.keep_together = True
    return paragraph


# ---------------------------------------------------------------------------
# 表格
# ---------------------------------------------------------------------------

def choose_table_widths(headers: list[str]) -> list[float]:
    count = len(headers)
    first = headers[0].strip() if headers else ""
    if count == 2:
        if first == "项目":
            return [3.2, 11.8]
        if first in ("问题", "现象", "阶段"):
            return [7.4, 7.6]
        return [6.4, 8.6]
    if count == 3:
        if first in ("层次", "层"):
            return [3.0, 3.8, 8.2]
        return [3.4, 3.8, 7.8]
    if count == 4:
        if "测试文件" in first:
            return [5.0, 3.2, 3.2, 3.6]
        return [3.4, 3.6, 4.0, 4.0]
    return [CONTENT_WIDTH_CM / count] * count


def _prepare_table_xml(table, widths: list[float]) -> None:
    tblpr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblBorders", "w:tblLayout"):
        for element in tblpr.findall(qn(tag)):
            tblpr.remove(element)
    set_table_width(table)
    set_table_borders(table)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)


def _add_cell_text(cell, text: str, *, header: bool = False,
                   bold: bool = False, center: bool = False) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.style = "TableText"
    paragraph.paragraph_format.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    )
    add_inline_text(
        paragraph,
        text,
        size=9,
        color=HEADER_COLOR if header else TEXT_COLOR,
        bold=header or bold,
    )


def add_table_block(doc: Document, headers: list[str],
                    rows: list[list[str]]) -> None:
    column_count = max(1, len(headers))
    widths = choose_table_widths(headers)
    table = doc.add_table(rows=1, cols=column_count)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _prepare_table_xml(table, widths)

    header_row = table.rows[0]
    header_row.height = Cm(0.8)
    header_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    repeat_header(header_row)
    prevent_row_split(header_row)
    for index in range(column_count):
        cell = header_row.cells[index]
        shade_cell(cell, "D9E2F3")
        text = headers[index] if index < len(headers) else ""
        _add_cell_text(cell, text, header=True, center=True)

    center_columns = {
        index for index, name in enumerate(headers)
        if any(key in name for key in ("用例数", "通过", "通过率", "序号", "数量"))
    }
    bold_first = len(headers) == 2 and headers[0].strip() == "项目"

    for row_index, values in enumerate(rows):
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        prevent_row_split(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            if row_index % 2 == 1:
                shade_cell(cell, "F7F9FC")
            text = values[column_index] if column_index < len(values) else ""
            _add_cell_text(
                cell,
                text,
                bold=bold_first and column_index == 0,
                center=column_index in center_columns,
            )

    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Cm(width)
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)


# ---------------------------------------------------------------------------
# Markdown 解析
# ---------------------------------------------------------------------------

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
UL_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
OL_RE = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
SEPARATOR_CELL_RE = re.compile(r"^:?-{3,}:?$")


def split_table_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]

    cells: list[str] = []
    buffer: list[str] = []
    in_code = False
    index = 0
    while index < len(text):
        char = text[index]
        if char == "\\" and index + 1 < len(text):
            buffer.append(char)
            buffer.append(text[index + 1])
            index += 2
            continue
        if char == "`":
            in_code = not in_code
            buffer.append(char)
        elif char == "|" and not in_code:
            cells.append("".join(buffer).strip())
            buffer = []
        else:
            buffer.append(char)
        index += 1
    cells.append("".join(buffer).strip())
    return cells


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    non_empty = [cell for cell in cells if cell]
    return bool(non_empty) and all(
        SEPARATOR_CELL_RE.match(cell) for cell in non_empty
    )


def is_table_start(lines: list[str], index: int) -> bool:
    return (
        index + 1 < len(lines)
        and lines[index].strip().startswith("|")
        and is_table_separator(lines[index + 1])
    )


def is_block_start(lines: list[str], index: int) -> bool:
    raw = lines[index]
    stripped = raw.strip()
    return (
        not stripped
        or stripped.startswith("```")
        or stripped.startswith("#")
        or stripped == "---"
        or is_table_start(lines, index)
        or bool(UL_RE.match(raw))
        or bool(OL_RE.match(raw))
    )


def smart_join(parts: list[str]) -> str:
    result = ""
    for part in parts:
        if not part:
            continue
        if (
            result
            and result[-1].isascii()
            and result[-1].isalnum()
            and part[0].isascii()
            and part[0].isalnum()
        ):
            result += " "
        result += part
    return result


def parse_blocks(lines: list[str]) -> list[tuple]:
    blocks: list[tuple] = []
    index = 0
    total = len(lines)
    while index < total:
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            index += 1
            code_lines: list[str] = []
            while index < total and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if index < total:
                index += 1
            while code_lines and not code_lines[-1].strip():
                code_lines.pop()
            blocks.append(("code", language, code_lines))
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            blocks.append((
                "heading",
                len(heading.group(1)),
                heading.group(2).strip(),
            ))
            index += 1
            continue

        if stripped == "---":
            index += 1
            continue

        if is_table_start(lines, index):
            headers = split_table_row(lines[index])
            index += 2
            rows: list[list[str]] = []
            while index < total and lines[index].strip().startswith("|"):
                row = split_table_row(lines[index])
                if any(cell for cell in row):
                    rows.append(row)
                index += 1
            blocks.append(("table", headers, rows))
            continue

        unordered = UL_RE.match(raw)
        if unordered:
            level = len(unordered.group(1)) // 2
            text = unordered.group(2).strip()
            index += 1
            while index < total and not is_block_start(lines, index):
                text = smart_join([text, lines[index].strip()])
                index += 1
            blocks.append(("list", level, False, "", text))
            continue

        ordered = OL_RE.match(raw)
        if ordered:
            level = len(ordered.group(1)) // 2
            marker = ordered.group(2)
            text = ordered.group(3).strip()
            index += 1
            while index < total and not is_block_start(lines, index):
                text = smart_join([text, lines[index].strip()])
                index += 1
            blocks.append(("list", level, True, marker, text))
            continue

        parts = [stripped]
        index += 1
        while index < total and not is_block_start(lines, index):
            parts.append(lines[index].strip())
            index += 1
        blocks.append(("paragraph", smart_join(parts)))
    return blocks


# ---------------------------------------------------------------------------
# 渲染与命令行入口
# ---------------------------------------------------------------------------

def render_blocks(doc: Document, blocks: list[tuple]) -> None:
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            add_heading_block(doc, block[1], block[2])
        elif kind == "paragraph":
            text = block[1]
            if text.startswith("姓名："):
                add_cover_name(doc, text)
            else:
                add_body_paragraph(doc, text)
        elif kind == "table":
            add_table_block(doc, block[1], block[2])
        elif kind == "list":
            add_list_item(doc, block[1], block[2], block[3], block[4])
        elif kind == "code":
            add_code_block(doc, block[2])


def convert(source: Path, output: Path) -> tuple[int, int, int]:
    text = source.read_text(encoding="utf-8")
    blocks = parse_blocks(text.splitlines())

    doc = Document()
    configure_document(doc)
    render_blocks(doc, blocks)

    props = doc.core_properties
    props.title = "实验报告（项目完成报告）——基于 MongoDB 的在线问卷系统"
    props.subject = "数据管理系统课程实验报告"
    props.author = "________"
    props.last_modified_by = "________"
    props.comments = "由 docs/项目完成报告.md 转换生成"
    props.keywords = "MongoDB; Flask; 在线问卷; 实验报告"
    props.category = "实验报告"

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

    headings = sum(1 for block in blocks if block[0] == "heading")
    tables = sum(1 for block in blocks if block[0] == "table")
    codes = sum(1 for block in blocks if block[0] == "code")
    return headings, tables, codes


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    source = (
        Path(sys.argv[1]) if len(sys.argv) > 1
        else root / "docs" / "项目完成报告.md"
    )
    output = (
        Path(sys.argv[2]) if len(sys.argv) > 2
        else root / "docs" / "项目完成报告.docx"
    )

    if not source.exists():
        print(f"源文件不存在：{source}", file=sys.stderr)
        return 1
    if not source.is_file():
        print(f"源路径不是文件：{source}", file=sys.stderr)
        return 1

    headings, tables, codes = convert(source, output)
    print(f"已生成：{output}")
    print(f"解析结果：{headings} 个标题、{tables} 个表格、{codes} 段代码")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())






